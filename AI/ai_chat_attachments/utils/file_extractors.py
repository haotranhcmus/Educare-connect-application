# -*- coding: utf-8 -*-
import base64
import io
import csv
import json
import logging
from odoo.exceptions import UserError

import PyPDF2
import openpyxl
from docx import Document

_logger = logging.getLogger(__name__)


class FileExtractors:
    """Service for extracting text from various file formats"""
    
    def __init__(self, env):
        self.env = env
    
    def extract_text(self, attachment):
        """
        Extract text from attachment based on mimetype
        
        Args:
            attachment: ir.attachment record
            
        Returns:
            str: Extracted text content or None
        """
        if not attachment:
            return None
        
        # Get file content
        try:
            file_content = base64.b64decode(attachment.datas)
        except Exception as e:
            _logger.exception(f"Failed to decode attachment {attachment.id}")
            raise UserError(f"Failed to decode file: {str(e)}")
        
        # Dispatch to appropriate extractor based on mimetype
        mimetype = attachment.mimetype or ''
        
        try:
            if mimetype == 'application/pdf':
                return self._extract_from_pdf(file_content)
            elif mimetype in ['application/vnd.ms-excel', 
                            'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet']:
                return self._extract_from_excel(file_content)
            elif mimetype in ['application/msword',
                            'application/vnd.openxmlformats-officedocument.wordprocessingml.document']:
                return self._extract_from_word(file_content)
            elif mimetype == 'text/csv':
                return self._extract_from_csv(file_content)
            elif mimetype and mimetype.startswith('text/'):
                return self._extract_from_text(file_content)
            elif mimetype == 'application/json':
                return self._extract_from_json(file_content)
            elif mimetype in ['application/xml', 'text/xml']:
                return self._extract_from_xml(file_content)
            else:
                _logger.warning(f"Unsupported mimetype for text extraction: {mimetype}")
                return None
                
        except Exception as e:
            _logger.exception(f"Text extraction failed for {attachment.name}")
            raise UserError(f"Failed to extract text from {attachment.name}: {str(e)}")
    
    def _extract_from_pdf(self, file_content):
        """Extract text from PDF using PyPDF2 < 3.0.0"""
        
        try:
            reader = PyPDF2.PdfFileReader(io.BytesIO(file_content))
            text_parts = []
            
            for page_num in range(reader.numPages):
                page = reader.getPage(page_num)
                text = page.extractText()
                
                if text and text.strip():
                    text_parts.append(text)
            
            full_text = "\n\n".join(text_parts) if text_parts else None
            
            if not full_text or not full_text.strip():
                return None
            
            return full_text
            
        except Exception as e:
            _logger.exception("PDF extraction failed")
            raise UserError(f"Failed to extract PDF text: {str(e)}")
    
    def _extract_from_excel(self, file_content):
        """Extract text from Excel (handles Pivot Tables)"""
        try:
            workbook = openpyxl.load_workbook(io.BytesIO(file_content), data_only=True)
            text_parts = []
            
            for sheet_name in workbook.sheetnames:
                sheet = workbook[sheet_name]
                text_parts.append(f"\n=== SHEET: {sheet_name} ===\n")
                
                # Get actual dimensions by checking all cells
                # This handles merged cells and Pivot Tables better
                max_row = 0
                max_col = 0
                
                for row in sheet.iter_rows():
                    for cell in row:
                        if cell.value is not None:
                            max_row = max(max_row, cell.row)
                            max_col = max(max_col, cell.column)
                
                if max_row == 0 or max_col == 0:
                    text_parts.append("(Empty sheet)")
                    continue
                
                # Extract ALL rows (including those with empty cells for Pivot Tables)
                rows = []
                for row_idx in range(1, max_row + 1):
                    row_values = []
                    has_content = False
                    
                    for col_idx in range(1, max_col + 1):
                        cell = sheet.cell(row_idx, col_idx)
                        cell_value = str(cell.value).strip() if cell.value is not None else ''
                        row_values.append(cell_value)
                        if cell_value:
                            has_content = True
                    
                    # Only skip COMPLETELY empty rows
                    if has_content:
                        rows.append(row_values)
                
                if not rows:
                    text_parts.append("(Empty sheet)")
                    continue
                
                text_parts.append(f"Total: {len(rows)} rows, {max_col} columns\n")
                
                for i, row in enumerate(rows, 1):
                    # Remove trailing empty cells
                    while row and not row[-1]:
                        row.pop()
                    
                    if row:
                        text_parts.append(f"Row {i}: " + " | ".join(row))
            
            return "\n".join(text_parts) if text_parts else None
            
        except Exception as e:
            _logger.exception("Excel extraction failed")
            raise UserError(f"Failed to extract Excel text: {str(e)}")
    
    def _extract_from_word(self, file_content):
        """Extract text from Word document"""
        try:
            doc = Document(io.BytesIO(file_content))
            text_parts = []
            
            for paragraph in doc.paragraphs:
                if paragraph.text.strip():
                    text_parts.append(paragraph.text)
            
            # Extract from tables
            for table in doc.tables:
                for row in table.rows:
                    row_text = [cell.text.strip() for cell in row.cells]
                    if any(row_text):
                        text_parts.append(" | ".join(row_text))
            
            return "\n".join(text_parts) if text_parts else None
            
        except Exception as e:
            _logger.exception("Word extraction failed")
            raise UserError(f"Failed to extract Word text: {str(e)}")
    
    def _extract_from_csv(self, file_content):
        """Extract text from CSV"""
        try:
            text = file_content.decode('utf-8')
            reader = csv.reader(io.StringIO(text))
            
            rows = []
            for row in reader:
                if any(cell.strip() for cell in row):
                    rows.append(" | ".join(row))
            
            return "\n".join(rows) if rows else None
            
        except Exception as e:
            _logger.exception("CSV extraction failed")
            raise UserError(f"Failed to extract CSV text: {str(e)}")
    
    def _extract_from_text(self, file_content):
        """Extract text from plain text file"""
        try:
            try:
                return file_content.decode('utf-8')
            except UnicodeDecodeError:
                return file_content.decode('latin-1')
        except Exception as e:
            _logger.exception("Text extraction failed")
            raise UserError(f"Failed to extract text: {str(e)}")
    
    def _extract_from_json(self, file_content):
        """Extract text from JSON file"""
        try:
            data = json.loads(file_content.decode('utf-8'))
            return json.dumps(data, indent=2, ensure_ascii=False)
        except Exception as e:
            _logger.exception("JSON extraction failed")
            raise UserError(f"Failed to extract JSON text: {str(e)}")
    
    def _extract_from_xml(self, file_content):
        """Extract text from XML file"""
        try:
            return file_content.decode('utf-8')
        except Exception as e:
            _logger.exception("XML extraction failed")
            raise UserError(f"Failed to extract XML text: {str(e)}")
